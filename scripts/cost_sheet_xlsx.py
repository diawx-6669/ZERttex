#!/usr/bin/env python3
"""Builds Cost_Payback_Sheet.xlsx with live formulas (not pasted values).
Input: a JSON dump of src/data/budget.ts (written by scripts/build-cost-sheet.cjs).
Usage: python3 scripts/cost_sheet_xlsx.py <config.json> <out.xlsx>
"""
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

cfg = json.load(open(sys.argv[1], encoding="utf-8"))
out = sys.argv[2]
m = cfg["model"]
c = cfg["caseFacts"]

wb = Workbook()
BOLD = Font(bold=True)
HEAD = PatternFill("solid", fgColor="12213F")
HEAD_FONT = Font(bold=True, color="FFFFFF")
INPUT = PatternFill("solid", fgColor="FFF2CC")  # editable assumptions
OUT = PatternFill("solid", fgColor="D9EAD3")  # key results
thin = Side(style="thin", color="BBBBBB")
BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
KZT = '#,##0" KZT"'
INT = "#,##0"
PCT = "0%"


def header(ws, row, values):
    for i, v in enumerate(values, 1):
        cell = ws.cell(row=row, column=i, value=v)
        cell.fill, cell.font = HEAD, HEAD_FONT


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ---------------- Assumptions ----------------
A = wb.active
A.title = "Assumptions"
A["A1"] = "Saryarqa Go · Cost and payback sheet · assumptions (yellow cells can be changed)"
A["A1"].font = Font(bold=True, size=13)
header(A, 3, ["Input", "Value", "Unit", "Label", "Note"])
rows = [
    ("Budget cap (development and launch, year 1)", cfg["cap"], "KZT", "Case fact", "Case rule", KZT, False),
    ("Travelers in 2025", c["travelers2025"], "people", "Case fact", "", INT, False),
    ("Revenue in 2025", c["revenue2025"], "KZT", "Case fact", "", KZT, False),
    ("Average price per traveler", m["avgPrice"], "KZT", "Case fact", "Stated in the case", KZT, False),
    ("Revenue / travelers (check)", "=B6/B5", "KZT", "Calculated", "44,118 vs the stated 44,000: the model uses 44,000 (careful)", KZT, False),
    ("Share of foreign travelers today", c["foreignShareToday"], "", "Case fact", "~96% are from Kazakhstan", PCT, False),
    ("Goal: share of foreign travelers", c["goalForeignShare"], "", "Case fact", "within 24 months", PCT, False),
    ("Contribution margin", m["margin"], "", "Assumption", "share of the price left after guides, transport, hotels, fees", PCT, True),
    ("Contribution per traveler", "=B7*B11", "KZT", "Calculated", "price × margin", KZT, False),
    ("Running cost per month from month 13", m["runningCostPerMonth"], "KZT", "Assumption", "hosting, AI usage, support, content", KZT, True),
    ("One-time costs spread over months 1..N", m["oneTimeMonths"], "months", "Assumption", "", INT, True),
    ("Other year-1 costs spread over months 1..N", m["otherMonths"], "months", "Assumption", "recurring + contingency", INT, True),
    ("Launch month", m["launchMonth"], "", "Assumption", "", INT, True),
    ("Month the ramp reaches 100%", m["rampEndMonth"], "", "Assumption", "linear ramp from launch, then flat", INT, True),
    ("Worse case: results × factor", m["downsideFactor"], "", "Assumption", "30% lower results", "0.00", True),
]
r = 4
for name, val, unit, label, note, fmt, editable in rows:
    A.cell(row=r, column=1, value=name)
    v = A.cell(row=r, column=2, value=val)
    v.number_format = fmt
    if editable:
        v.fill = INPUT
    A.cell(row=r, column=3, value=unit)
    A.cell(row=r, column=4, value=label)
    A.cell(row=r, column=5, value=note)
    r += 1
# Rows: B4 cap, B5 travelers, B6 revenue, B7 price, B8 rev/trav, B9 share today, B10 goal, B11 margin,
# B12 contribution, B13 running, B14 one-time months, B15 other months, B16 launch, B17 ramp end, B18 downside

A["A20"] = "Where the extra foreign travelers come from"
A["A20"].font = BOLD
header(A, 21, ["Step", "Value", "", "Label", "Formula"])
deriv = [
    ("Foreign travelers today", "=ROUND(B5*B9,0)", "Calculated", "3,400 × 4%"),
    ("Domestic travelers (kept)", "=B5-B22", "Calculated", "3,400 − foreign"),
    ("Total needed for 30% foreign", "=ROUND(B23/(1-B10),0)", "Calculated", "domestic ÷ 0.7"),
    ("Foreign travelers needed", "=ROUND(B24*B10,0)", "Calculated", "total × 30%"),
    ("Extra foreign travelers per year (target by month 24)", "=B25-B22", "Calculated", "needed − today ≈ 1,263"),
]
r = 22
for name, f, label, note in deriv:
    A.cell(row=r, column=1, value=name)
    A.cell(row=r, column=2, value=f).number_format = INT
    A.cell(row=r, column=4, value=label)
    A.cell(row=r, column=5, value=note)
    r += 1
widths(A, [52, 16, 9, 13, 62])

# ---------------- Budget ----------------
B = wb.create_sheet("Budget")
B["A1"] = "Year-1 budget (planning assumptions, team proposal)"
B["A1"].font = Font(bold=True, size=13)
header(B, 3, ["Item", "Type", "Gate", "KZT", "Share"])
first = 4
for i, line in enumerate(cfg["lines"]):
    rr = first + i
    B.cell(row=rr, column=1, value=line["item"])
    B.cell(row=rr, column=2, value=line["kind"])
    B.cell(row=rr, column=3, value=line["gate"])
    amt = B.cell(row=rr, column=4, value=line["amount"])
    amt.number_format, amt.fill = INT, INPUT
    B.cell(row=rr, column=5, value=f"=D{rr}/$D${first + len(cfg['lines'])}").number_format = "0.0%"
last = first + len(cfg["lines"]) - 1
tot = last + 1
B.cell(row=tot, column=1, value="Total year 1").font = BOLD
B.cell(row=tot, column=4, value=f"=SUM(D{first}:D{last})").number_format = INT
B.cell(row=tot, column=4).font = BOLD
B.cell(row=tot, column=5, value=f"=SUM(E{first}:E{last})").number_format = "0.0%"
summary = [
    ("One-time costs", f'=SUMIF(B{first}:B{last},"one-time",D{first}:D{last})'),
    ("Recurring + contingency (year 1)", f"=D{tot}-D{tot + 2}"),
    ("Monthly cost, months 1-3", f"=D{tot + 2}/Assumptions!B14+D{tot + 3}/Assumptions!B15"),
    ("Monthly cost, months 4-12", f"=D{tot + 3}/Assumptions!B15"),
    ("Monthly cost from month 13", "=Assumptions!B13"),
    ("Gate 1: MVP + pilot", f"=SUMIF(C{first}:C{last},1,D{first}:D{last})"),
    ("Gate 2: only if pilot metrics are met", f"=SUMIF(C{first}:C{last},2,D{first}:D{last})"),
    ("Within the cap?", f'=IF(D{tot}<=Assumptions!B4,"YES","NO")'),
]
for i, (name, f) in enumerate(summary):
    rr = tot + 2 + i
    B.cell(row=rr, column=1, value=name)
    B.cell(row=rr, column=4, value=f).number_format = INT
B.cell(row=tot + 11, column=1, value="AI usage is paid from the existing line 'Hosting, AI usage, maintenance'. Staged spending (Gate 1 / Gate 2) is a team proposal.").font = Font(italic=True, color="666666")
widths(B, [44, 12, 7, 16, 9])
ONE = f"Budget!D{tot + 2}"
OTHER = f"Budget!D{tot + 3}"

# ---------------- Payback ----------------
P = wb.create_sheet("Payback")
P["A1"] = "Monthly cash flow, 48 months (base and worse case −30%)"
P["A1"].font = Font(bold=True, size=13)
header(P, 3, ["Month", "Costs", "Ramp", "Travelers (base)", "Margin earned (base)", "Net (base)", "Cumulative (base)", "Travelers (−30%)", "Margin earned (−30%)", "Cumulative (−30%)", "Payback? base", "Payback? −30%"])
start = 4
H = m["horizonMonths"]
for i in range(H):
    rr = start + i
    mo = i + 1
    P.cell(row=rr, column=1, value=mo)
    P.cell(row=rr, column=2, value=f"=IF(A{rr}<=Assumptions!$B$14,{ONE}/Assumptions!$B$14,0)+IF(A{rr}<=Assumptions!$B$15,{OTHER}/Assumptions!$B$15,Assumptions!$B$13)").number_format = INT
    P.cell(row=rr, column=3, value=f"=IF(A{rr}<Assumptions!$B$16,0,IF(A{rr}>=Assumptions!$B$17,1,(A{rr}-Assumptions!$B$16+1)/(Assumptions!$B$17-Assumptions!$B$16+1)))").number_format = "0.000"
    P.cell(row=rr, column=4, value=f"=Assumptions!$B$26/12*C{rr}").number_format = "0.0"
    P.cell(row=rr, column=5, value=f"=D{rr}*Assumptions!$B$12").number_format = INT
    P.cell(row=rr, column=6, value=f"=E{rr}-B{rr}").number_format = INT
    P.cell(row=rr, column=7, value=f"=F{rr}" if mo == 1 else f"=G{rr - 1}+F{rr}").number_format = INT
    P.cell(row=rr, column=8, value=f"=D{rr}*Assumptions!$B$18").number_format = "0.0"
    P.cell(row=rr, column=9, value=f"=H{rr}*Assumptions!$B$12").number_format = INT
    P.cell(row=rr, column=10, value=f"=I{rr}-B{rr}" if mo == 1 else f"=J{rr - 1}+I{rr}-B{rr}").number_format = INT
    P.cell(row=rr, column=11, value=f'=IF(AND(A{rr}>=Assumptions!$B$16,G{rr}>=0),A{rr},"")')
    P.cell(row=rr, column=12, value=f'=IF(AND(A{rr}>=Assumptions!$B$16,J{rr}>=0),A{rr},"")')
end = start + H - 1
widths(P, [8, 13, 8, 15, 18, 13, 17, 16, 19, 17, 13, 13])
P.freeze_panes = "B4"

# ---------------- Summary ----------------
S = wb.create_sheet("Summary", 0)
S["A1"] = "Saryarqa Go · Cost and payback summary (all values are formulas)"
S["A1"].font = Font(bold=True, size=13)
header(S, 3, ["Result", "Base", "Worse case (−30%)", "Label"])
res = [
    ("Total year-1 budget (KZT)", f"=Budget!D{tot}", f"=Budget!D{tot}", "Proposal", INT),
    ("Budget cap (KZT)", "=Assumptions!B4", "=Assumptions!B4", "Case fact", INT),
    ("Contribution per traveler (KZT)", "=Assumptions!B12", "=Assumptions!B12", "Calculated", INT),
    ("Extra foreign travelers per year at month 24", "=Assumptions!B26", "=Assumptions!B26*Assumptions!B18", "Calculated", INT),
    ("Break-even travelers (cap ÷ contribution)", "=ROUNDUP(Assumptions!B4/Assumptions!B12,0)", "=ROUNDUP(Assumptions!B4/Assumptions!B12,0)", "Calculated", INT),
    ("Lowest cumulative cash (KZT)", f"=MIN(Payback!G{start}:G{end})", f"=MIN(Payback!J{start}:J{end})", "Calculated", INT),
    ("Cumulative cash at month 24 (KZT)", f"=Payback!G{start + 23}", f"=Payback!J{start + 23}", "Calculated", INT),
    ("Cumulative cash at month 48 (KZT)", f"=Payback!G{end}", f"=Payback!J{end}", "Calculated", INT),
    ("PAYBACK MONTH", f'=IF(COUNT(Payback!K{start}:K{end})=0,"> 48",MIN(Payback!K{start}:K{end}))', f'=IF(COUNT(Payback!L{start}:L{end})=0,"> 48",MIN(Payback!L{start}:L{end}))', "Calculated", "0"),
]
for i, (name, b, d, label, fmt) in enumerate(res):
    rr = 4 + i
    S.cell(row=rr, column=1, value=name)
    for col, f in ((2, b), (3, d)):
        cell = S.cell(row=rr, column=col, value=f)
        cell.number_format = fmt
        cell.border = BOX
        if name == "PAYBACK MONTH":
            cell.fill, cell.font = OUT, BOLD
    S.cell(row=rr, column=4, value=label)
S["A14"] = "Payback month = first month (from launch) when cumulative cash is ≥ 0. Worse case = 30% fewer extra travelers, same costs."
S["A15"] = "Check: 150,000,000 ÷ 3,400 = 44,118 KZT, slightly above the stated 44,000 average price. The model uses 44,000 (careful)."
S["A16"] = "Planning assumptions, not a forecast. Change the yellow cells on 'Assumptions' and 'Budget'; every result updates."
for a in ("A14", "A15", "A16"):
    S[a].font = Font(italic=True, color="555555")
widths(S, [48, 16, 19, 12])

wb.save(out)
print("written", out)
